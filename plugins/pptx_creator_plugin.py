"""PowerPoint Creator Plugin - Create .pptx presentations."""

import logging
from pathlib import Path

logger = logging.getLogger("jiro.plugins.pptx_creator")

PLUGIN_NAME = "pptx_creator"
PLUGIN_DESCRIPTION = "Create PowerPoint presentations (.pptx)"
PLUGIN_COMMANDS = ["create ppt", "create presentation", "make ppt", "new presentation", "create pptx"]


async def handle(command: str, context: dict = None) -> str:
    """Create a PowerPoint presentation."""
    try:
        from pptx import Presentation
        from pptx.util import Inches
    except ImportError:
        return "python-pptx is not installed. Run: pip install python-pptx"

    lower = command.lower()
    for prefix in PLUGIN_COMMANDS:
        if lower.startswith(prefix):
            rest = command[len(prefix):].strip()
            break
    else:
        rest = command

    parts = rest.split(maxsplit=1)
    if not parts:
        return "Please specify a title. Example: 'create ppt MyPresentation Slide content here'"

    title = parts[0]
    filename = title
    if not filename.endswith('.pptx'):
        filename += '.pptx'

    content = parts[1] if len(parts) > 1 else ""

    desktop = Path.home() / "Desktop"
    if desktop.exists():
        filepath = desktop / filename
    else:
        docs_dir = Path.home() / "Documents"
        docs_dir.mkdir(exist_ok=True)
        filepath = docs_dir / filename

    try:
        prs = Presentation()

        # Title slide
        slide_layout = prs.slide_layouts[0]
        slide = prs.slides.add_slide(slide_layout)
        slide.shapes.title.text = title.replace('.pptx', '')
        slide.placeholders[1].text = "Created by Jiro AI"

        # Content slides
        if content:
            slides_text = content.split('|')
            for slide_content in slides_text:
                slide_layout = prs.slide_layouts[1]
                slide = prs.slides.add_slide(slide_layout)
                parts = slide_content.strip().split(':', 1)
                if len(parts) == 2:
                    slide.shapes.title.text = parts[0].strip()
                    slide.placeholders[1].text = parts[1].strip()
                else:
                    slide.shapes.title.text = "Slide"
                    slide.placeholders[1].text = slide_content.strip()

        prs.save(str(filepath))
        return f"Presentation created: {filepath}"
    except Exception as e:
        logger.error("PPT creation failed: %s", e)
        return f"Failed to create presentation: {e}"
